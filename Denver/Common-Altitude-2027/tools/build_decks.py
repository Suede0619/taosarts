"""Create original editable PowerPoint decks, slide scripts and brand vectors."""
from pathlib import Path
import json, math, re
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
BASE=Path(__file__).resolve().parents[1]
DATA=BASE/"data"
P=json.loads((BASE/"data/project.json").read_text())
F=json.loads((BASE/"data/financial-results.json").read_text())
OUT=BASE/"decks";OUT.mkdir(exist_ok=True)
ASSETS=BASE/"assets";ASSETS.mkdir(exist_ok=True)
C=P["palette"]
FONT_DIR=Path("/System/Library/Fonts/Supplemental")
FONT_NORMAL=FONT_DIR/"Arial.ttf";FONT_BOLD=FONT_DIR/"Arial Bold.ttf"
USD=lambda x:("$"+f"{x:,.0f}")
T={"core_rooms":str(F["core"]["rooms"]),"core_revenue":USD(F["core"]["organizer_revenue"]),
   "core_cost":USD(F["core"]["cost"]),"core_surplus":USD(F["core"]["pre_owner_surplus"]),
   "eligible_pool":USD(F["waterfall"]["initial_eligible_pool"]),"initial_owner":USD(F["waterfall"]["initial_owner_share"]),
   "final_owner":USD(F["waterfall"]["owner_total_if_hold_released_without_claims"]),
   "cash_need":USD(F["waterfall"]["peak_bridge_with_sponsor_hold"])}
def subst(x):
    if isinstance(x,str):
        for k,v in T.items():x=x.replace("{{"+k+"}}",v)
        assert "{{" not in x,x
        return x
    if isinstance(x,list):return [subst(v) for v in x]
    if isinstance(x,dict):return {k:subst(v) for k,v in x.items()}
    return x
def rgb(h):return RGBColor.from_string(h.strip("#"))
def rect(sl,x,y,w,h,color,line=None,width=1):
    sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(color)
    if line:sh.line.color.rgb=rgb(line);sh.line.width=Pt(width)
    else:sh.line.fill.background()
    return sh
def measure_wrap(text,width,size,bold=False):
    f=ImageFont.truetype(str(FONT_BOLD if bold else FONT_NORMAL),int(size*4))
    result=[]
    for para in str(text).split("\n"):
        if not para:result.append("");continue
        words=para.split();line=""
        for word in words:
            candidate=(line+" "+word).strip()
            if f.getlength(candidate)/4<=width*72:line=candidate
            else:
                if line:result.append(line)
                line=word
                if f.getlength(line)/4>width*72:return None
        result.append(line)
    return result
def text(sl,x,y,w,h,value,size=22,bold=False,color=None,family="Arial",minsize=16,leading=1.18):
    chosen=None
    for fs in range(int(size),int(minsize)-1,-1):
        lines=measure_wrap(value,w,fs,bold)
        if lines is not None and len(lines)*fs*leading<=h*72:
            chosen=(fs,lines);break
    if chosen is None:raise ValueError(f"Text overflow: {value!r} in {w}x{h}")
    fs,lines=chosen
    box=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame;tf.clear();tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=family;p.font.size=Pt(fs);p.font.bold=bold
        p.font.color.rgb=rgb(color or C["ink"]);p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=1.08
    return box,len(lines)*fs*leading/72
def mark(sl,x,y,w,h,color):
    actual_w=min(w,h/1.35);actual_h=actual_w*1.35
    x+=(w-actual_w)/2;y+=(h-actual_h)/2;w=actual_w;h=actual_h
    stroke=max(.035,w*.028)
    rect(sl,x,y,w*.72,stroke,color)
    rect(sl,x,y,stroke,h*.8,color)
    rect(sl,x+w*.72-stroke,y,stroke,h*.52,color)
    rect(sl,x+w*.28,y+h*.2,w*.72,stroke,color)
    rect(sl,x+w*.28,y+h*.2,stroke,h*.8,color)
    rect(sl,x+w-stroke,y+h*.2,stroke,h*.8,color)
    rect(sl,x+w*.28,y+h-stroke,w*.72,stroke,color)
    rect(sl,x,y+h*.8-stroke,w*.48,stroke,color)
def footer(sl,i,total,dark):
    color=C["paper"] if dark else C["ink"]
    rect(sl,1,8.35,14,.012,C["stone"])
    text(sl,1,8.50,4,.23,"COMMON ALTITUDE / 2027",10,True,color,minsize=10)
    text(sl,5,8.50,7,.23,"WORKING PROPOSAL / 24 OCT 2026",9,False,color,minsize=9)
    text(sl,14.15,8.50,.85,.23,f"{i:02d} / {total:02d}",9,False,color,minsize=9)
def header(sl,d,deck):
    text(sl,1,.45,13.6,.3,deck["audience"].upper(),10,True,C["stone"],minsize=9)
    _,th=text(sl,1,1.08,14,2.35,d["title"],50,True,minsize=40,leading=1.1)
    y=1.08+th+.22
    _,sh=text(sl,1,y,13.8,1.1,d["subtitle"],19,minsize=17)
    return max(3.12,y+sh+.35)
def build_slide(prs,d,deck,index,total):
    sl=prs.slides.add_slide(prs.slide_layouts[6])
    kind=d["layout"];dark=kind in ["cover","ask"]
    bg=sl.background.fill;bg.solid();bg.fore_color.rgb=rgb(C["ink"] if dark else C["paper"])
    if kind=="cover":
        text(sl,1,.6,13.5,.35,deck["audience"].upper(),11,True,C["citron"],minsize=10)
        text(sl,1,1.55,9.5,3.2,d["title"],80,True,C["paper"],minsize=70,leading=1.08)
        text(sl,1,5.25,8.9,1.65,d["subtitle"],23,False,C["paper"],minsize=20)
        mark(sl,11,1.65,3.7,5.3,C["citron"])
        text(sl,10.9,7.3,4,.45,"ORIGINAL CONCEPT GRAPHIC",9,True,C["stone"],minsize=9)
    elif kind=="ask":
        text(sl,1,.6,13.5,.35,"THE NEXT DECISION",11,True,C["citron"],minsize=10)
        _,h=text(sl,1,1.5,10.3,2.2,d["title"],52,True,C["paper"],minsize=42)
        y=1.5+h+.35
        _,h2=text(sl,1,y,10.5,1.5,d["subtitle"],23,False,C["paper"],minsize=19)
        by=max(4.55,y+h2+.35)
        for i,item in enumerate(d["items"]):
            rect(sl,1,by+i*.8,.08,.42,C["citron"])
            text(sl,1.3,by+i*.8,10.1,.7,item,20,False,C["paper"],minsize=17)
        mark(sl,12.4,4.3,2.3,3.1,C["citron"])
    else:
        by=header(sl,d,deck);avail=8.03-by
        if kind in ["trio","week","timeline","packages","metrics"]:
            items=d["items"];n=len(items);gap=.28;w=(14-gap*(n-1))/n
            for i,(a,b) in enumerate(items):
                x=1+i*(w+gap)
                rect(sl,x,by,w,avail,C["ink"] if kind=="metrics" else "#E7E5DE")
                rect(sl,x,by,w,.065,C["citron"] if i%2==0 else C["cobalt"])
                if kind=="metrics":
                    text(sl,x+.28,by+.4,w-.56,1.1,a,42,True,C["citron"],minsize=29)
                    text(sl,x+.28,by+1.85,w-.56,avail-2.05,b,21,False,C["paper"],minsize=18)
                else:
                    text(sl,x+.28,by+.42,w-.56,1.3,a,27,True,minsize=21)
                    text(sl,x+.28,by+1.92,w-.56,avail-2.15,b,21,minsize=17)
        elif kind=="waterfall":
            items=d["items"];rowh=avail/len(items)
            for i,(a,b) in enumerate(items):
                rect(sl,1,by+i*rowh,14,rowh-.13,C["ink"] if i==len(items)-1 else "#E7E5DE")
                ink=C["paper"] if i==len(items)-1 else C["ink"]
                text(sl,1.3,by+i*rowh+.18,3.9,rowh-.25,a,30,True,C["citron"] if i==len(items)-1 else ink,minsize=24)
                text(sl,5.5,by+i*rowh+.18,9.1,rowh-.25,b,19,False,ink,minsize=17)
        elif kind=="rooms":
            gx,gy=1,by;gw,gh=7.5,avail-.4
            cw=gw/10;ch=gh/6
            for row in range(6):
                for col in range(10):
                    cell_index=row*10+col
                    fill=C["ink"] if cell_index<24 else (C["citron"] if cell_index<36 else (C["cobalt"] if cell_index<42 else "#D8DAD3"))
                    rect(sl,gx+col*cw,gy+row*ch,cw-.08,ch-.08,fill)
            for i,(a,b) in enumerate(d["items"]):
                y=by+i*(avail/4)
                text(sl,9.2,y,1.2,.7,a,33,True,minsize=29)
                text(sl,10.7,y+.08,4.15,.8,b,20,minsize=17)
            text(sl,1,7.93,7.8,.24,"CONCEPTUAL FUNCTION GRID / NOT A SURVEYED FLOOR PLAN",9,True,minsize=9)
        elif kind=="palette":
            n=5;w=14/n
            for i,(name,color) in enumerate(C.items()):
                rect(sl,1+i*w,by,w-.10,avail-1.1,color)
                text(sl,1+i*w,by+avail-.85,w-.1,.35,name.upper(),14,True,minsize=12)
                text(sl,1+i*w,by+avail-.42,w-.1,.3,color,13,minsize=12)
        elif kind=="mark":
            mark(sl,1.4,by,4.4,avail-.35,C["ink"])
            for i,item in enumerate(d["items"]):text(sl,8,by+i*1.05,6.6,.85,item,22,minsize=18)
        elif kind=="poster":
            rect(sl,1,by,5.2,avail,C["ink"])
            text(sl,1.3,by+.3,4.6,1.8,"ART IS\nTAKING\nOVER.",31,True,C["citron"],minsize=27)
            text(sl,1.3,by+2.5,4.5,.8,"COMMON ALTITUDE\nDENVER / 2027",17,True,C["paper"],minsize=15)
            for i,item in enumerate(d["items"]):text(sl,7,by+i*1.05,7.7,.9,item,22,minsize=18)
        else:
            items=d["items"];n=max(1,len(items));rowh=avail/n
            for i,item in enumerate(items):
                rect(sl,1,by+i*rowh+.03,.075,min(.5,rowh-.15),C["cobalt"] if kind=="sources" else C["ink"])
                box,_=text(sl,1.35,by+i*rowh,13.4,rowh-.13,item,23 if kind!="sources" else 18,minsize=17)
                if kind=="sources":
                    box.click_action.hyperlink.address=("https://www.denverenergycenter.com/home" if "denverenergycenter.com" in item else "https://github.com/Suede0619/taosarts/blob/main/Denver/Common-Altitude-2027/RESEARCH.md")
    footer(sl,index,total,dark)
    note=sl.notes_slide.notes_text_frame
    note.text=(deck["audience"]+"\n"+d["title"].replace("\n"," ")+"\n\n"+d["subtitle"]+
               "\n\n"+ "\n".join((" / ".join(x) if isinstance(x,list) else x) for x in d["items"])+
               "\n\nStatus: working proposal; site/date/partners/funding/capacities unconfirmed. No external action is authorized by the deck.\n"+
               "Financial source: data/project.json and generated financial-results.json.\n"+
               "Primary venue reference: https://www.denverenergycenter.com/home\n"+
               "Program/market/rights sources and individual candidate evidence: RESEARCH.md in the accompanying package.")
    return sl

def assets():
    svg='<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 240 324"><title>Common Altitude original abstract mark</title><path d="M20 243V27H164V167.4 M76 297V81H220V297H76 M20 243H116" fill="none" stroke="#151A22" stroke-width="6" stroke-linejoin="miter"/></svg>'
    (ASSETS/"common-altitude-mark.svg").write_text(svg+"\n")
    word='<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1100 340"><title>Common Altitude working wordmark</title><rect width="1100" height="340" fill="#F3F0E9"/><g transform="translate(30,5)">'+svg.replace('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 240 324">','<g>').replace('</svg>','</g>')+'</g><text x="315" y="145" font-family="Arial,sans-serif" font-weight="700" font-size="74" fill="#151A22">COMMON</text><text x="315" y="235" font-family="Arial,sans-serif" font-weight="700" font-size="74" fill="#151A22">ALTITUDE</text></svg>'
    (ASSETS/"common-altitude-wordmark.svg").write_text(word+"\n")
    poster='<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1200 1800"><title>Concept poster, dates proposed</title><rect width="1200" height="1800" fill="#151A22"/><text x="90" y="155" fill="#DCFF5F" font-family="Arial,sans-serif" font-size="33" font-weight="700">COMMON ALTITUDE / DENVER</text><text x="80" y="440" fill="#DCFF5F" font-family="Arial,sans-serif" font-size="160" font-weight="700"><tspan x="80">ART IS</tspan><tspan x="80" dy="180">TAKING</tspan><tspan x="80" dy="180">OVER.</tspan></text><path d="M100 1280V950H740V1170M360 1460V1070H1100V1460H360M100 1280H600" fill="none" stroke="#F3F0E9" stroke-width="16"/><text x="90" y="1580" fill="#F3F0E9" font-family="Arial,sans-serif" font-size="37">OFFICE OF THE IMPOSSIBLE</text><text x="90" y="1650" fill="#F3F0E9" font-family="Arial,sans-serif" font-size="28">Proposed August 9–15, 2027 / Art Fair &amp; Exhibition</text><text x="90" y="1730" fill="#A6AAA6" font-family="Arial,sans-serif" font-size="20">CONCEPT FOR REVIEW / VENUE, DATE AND PARTICIPANTS UNCONFIRMED</text></svg>'
    (ASSETS/"concept-poster.svg").write_text(poster+"\n")

def main():
    decks=subst(json.loads((BASE/"data/decks.json").read_text()))
    brand={"id":"00-Creative-Direction","audience":"Creative direction / working brand system","title":"A precise interruption.","slides":[
       {"layout":"cover","title":"COMMON\nALTITUDE","subtitle":"OFFICE OF THE IMPOSSIBLE\nWorking creative direction / Denver 2027","items":[]},
       {"layout":"statement","title":"The ordinary becomes\nextraordinary.","subtitle":"A remarkable place to discover and acquire art, open to the city.","items":["Artist authorship, a public welcome and a serious art market.","A familiar room can become an unexpected encounter.","The visitor experience should be as good as the headline idea."]},
       {"layout":"palette","title":"A precise interruption.","subtitle":"Ink and paper establish calm. Citron and cobalt create a controlled moment of surprise.","items":[]},
       {"layout":"mark","title":"An open frame.\nA shared passage.","subtitle":"An original abstract identity device, not a drawing of the venue.","items":["A room, a threshold and a horizon.","Two structures in one coherent system.","Simple enough for a label; strong enough for a facade-scale concept.","Final name/mark clearance and production proofs remain required."]},
       {"layout":"poster","title":"A clear invitation.","subtitle":"A concept application with honest proposed status, not a public event announcement.","items":["Oversized short language.","A disciplined original graphic.","Practical information remains readable.","Artwork is credited evidence, never decorative stock."]},
       {"layout":"trio","title":"One system. Many voices.","subtitle":"The art leads; the system supplies coherence.","items":[["PRECISE","Editorial clarity, accurate information and serious service."],["NOCTURNAL","Artist-led light/sound for specific approved programs."],["MATERIAL","Hands, making and paid local authorship."]]},
       {"layout":"ask","title":"Make the experience\nworthy of the art.","subtitle":"Review the direction with artists, buyers, residents and disabled visitors; refine through funded design work.","items":["Keep authorship, rights and public welcome intact.","Choose achievable exceptional work over borrowed spectacle.","Use the brand book and production briefs as the working system."]}
    ]}
    decks=[brand]+decks
    manifest=[]
    for deck in decks:
        prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
        prs.core_properties.title="COMMON ALTITUDE — "+deck["audience"]
        prs.core_properties.author="Common Altitude planning"
        prs.core_properties.last_modified_by="Common Altitude planning"
        prs.core_properties.subject="Conditional proposal for review, October 24 2026"
        for i,d in enumerate(deck["slides"],1):build_slide(prs,d,deck,i,len(deck["slides"]))
        dest=OUT/(deck["id"]+".pptx");prs.save(dest)
        md=[f'# {deck["audience"]}',"","Working designed proposal · October 24, 2026. No site/date, partner, funding or capacity commitment. This slide script accompanies the editable PowerPoint and PDF.",""]
        for i,d in enumerate(deck["slides"],1):
            md += [f'## {i}. '+d["title"].replace("\n"," "),"",d["subtitle"],""]
            for item in d["items"]:md.append("- "+(" — ".join(item) if isinstance(item,list) else item))
            md.append("")
        (OUT/(deck["id"]+".md")).write_text("\n".join(md).rstrip()+"\n")
        manifest.append({"id":deck["id"],"audience":deck["audience"],"slides":len(deck["slides"]),
                         "pptx":dest.name,"pdf":deck["id"]+".pdf","script":deck["id"]+".md",
                         "status":"Designed working proposal; PDF export/visual QA required"})
    assets()
    (DATA/"deck-manifest.json").write_text(json.dumps(manifest,indent=2)+"\n")
    print(json.dumps({"presentations":len(manifest),"pitch_decks":len(decks)-1,"slides":sum(x["slides"] for x in manifest)}))
if __name__=="__main__":main()
